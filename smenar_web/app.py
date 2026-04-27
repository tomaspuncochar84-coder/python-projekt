from flask import Flask, render_template, request, jsonify, send_file
import csv
import os
import shutil
from generator import (
    generate_schedule, generate_actual_schedule,
    get_easter_dates, format_minutes, is_holiday
)
from datetime import datetime, timedelta

app = Flask(__name__)

UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

CZECH_HOLIDAYS = [
    "01.01.", "01.05.", "08.05.",
    "05.07.", "06.07.",
    "28.09.", "28.10.",
    "17.11.",
    "24.12.", "25.12.", "26.12."
]


def get_shift_time(shift, date_obj):
    is_tuesday = date_obj.weekday() == 1
    if shift == "Ran":
        return "4:00–12:01" if is_tuesday else "4:24–12:01"
    if shift == "Mol":
        if date_obj.weekday() == 5:
            return "5:24–18:15"
        elif date_obj.weekday() == 6:
            return "7:21–20:21"
        else:
            return "ručně"
    mapping = {
        "Odp": "11:46–20:21",
        "N":   "Nemoc",
        "Den": "5:24–18:21",
        "D":   "Dovolená",
        "Vol": ""
    }
    return mapping.get(shift, "")


@app.route("/", methods=["GET", "POST"])
def index():
    return render_template("index.html")


@app.route("/update_cell", methods=["POST"])
def update_cell():
    data = request.json
    datum = data.get("Datum")
    zamestnanec = data.get("Zaměstnanec")
    actual_shift = data.get("ActualShift", "")
    minuty_raw = data.get("Minuty", "")

    rows = []
    try:
        with open("overrides.csv", "r", encoding="utf-8") as f:
            reader = csv.DictReader(f, delimiter=";")
            rows = list(reader)
    except FileNotFoundError:
        pass

    def convert_to_minutes(val):
        if val is None or str(val).strip() == "":
            return None
        val = str(val).strip()
        if ":" in val:
            try:
                h, m = val.split(":")
                return int(h) * 60 + int(m)
            except:
                return None
        try:
            return int(val)
        except:
            return None

    novy_vykon = convert_to_minutes(minuty_raw)

    # "0" = signál pro odemknutí → smažeme ruční minuty
    unlocking = (str(minuty_raw).strip() == "0")

    found = False
    for row in rows:
        if row["Datum"] == datum and row["Zaměstnanec"] == zamestnanec:
            if actual_shift:
                row["Override"] = actual_shift
            if unlocking:
                row["Minuty"] = ""   # vymazání ručního zápisu = odemknutí
            elif novy_vykon is not None:
                row["Minuty"] = str(novy_vykon)
            found = True
            break

    if not found:
        rows.append({
            "Datum": datum,
            "Zaměstnanec": zamestnanec,
            "Override": actual_shift,
            "Minuty": "" if unlocking else (str(novy_vykon) if novy_vykon is not None else "")
        })

    with open("overrides.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["Datum", "Zaměstnanec", "Override", "Minuty"],
            delimiter=";"
        )
        writer.writeheader()
        writer.writerows(rows)

    substitutes, totals_by_month, employees_by_month = generate_actual_schedule()

    return jsonify({
        "employees_by_month": employees_by_month,
        "totals": totals_by_month,
        "substitutes": substitutes
    })


@app.route("/get_data")
def get_data():
    substitutes, totals_by_month, employees_by_month = generate_actual_schedule()
    return jsonify({
        "employees_by_month": employees_by_month,
        "totals": totals_by_month,
        "substitutes": substitutes
    })


@app.route("/regenerate", methods=["POST"])
def regenerate():
    from generator import generate_schedule
    start = datetime(2026, 1, 1)
    end   = datetime(2026, 12, 31)
    generate_schedule(start, end)
    substitutes, totals_by_month, employees_by_month = generate_actual_schedule()
    return jsonify({"status": "ok"})


@app.route("/export")
def export():
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    wb.remove(wb.active)

    with open("smenar_actual.csv", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        data = list(reader)

    months = {}
    for row in data:
        date = datetime.strptime(row["Datum"], "%d.%m.%Y")
        month = date.strftime("%m")
        if month not in months:
            months[month] = []
        months[month].append(row)

    month_names = {
        "01": "Leden", "02": "Únor", "03": "Březen",
        "04": "Duben", "05": "Květen", "06": "Červen",
        "07": "Červenec", "08": "Srpen", "09": "Září",
        "10": "Říjen", "11": "Listopad", "12": "Prosinec"
    }

    weekend_fill  = PatternFill(start_color="00CCFF", end_color="00CCFF", fill_type="solid")
    holiday_fill  = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    vacation_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    manual_fill   = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for month, rows in months.items():
        ws = wb.create_sheet(title=month_names[month])

        headers = ["Datum", "Zaměstnanec", "Plán", "Actual", "Výkon", "Dovolená", "Nemoc"]
        ws.append(headers)

        for row in rows:
            ws.append([
                row["Datum"],
                row["Zaměstnanec"],
                row["PlannedShift"],
                row["ActualShift"],
                format_minutes(int(row["Výkon"])) if row["Výkon"] else "",
                format_minutes(int(row["Dovolená"])) if row["Dovolená"] else "",
                format_minutes(int(row["Nemoc"])) if row["Nemoc"] else "",
            ])

        # Auto šířka
        for col in ws.columns:
            max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
            ws.column_dimensions[col[0].column_letter].width = max_length + 2

        # Bordery
        for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row_cells:
                cell.border = border

        # Tabulka
        table = Table(
            displayName=f"Table_{month}",
            ref=f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}"
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(table)

        # Barvy
        for i, row in enumerate(rows, start=2):
            date_obj = datetime.strptime(row["Datum"], "%d.%m.%Y")
            fill = None
            if is_holiday(date_obj):
                fill = holiday_fill
            elif date_obj.weekday() >= 5:
                fill = weekend_fill
            elif row["ActualShift"] == "D":
                fill = vacation_fill
            elif row.get("ManualMinutes") == "1":
                fill = manual_fill

            if fill:
                for col in range(1, len(headers) + 1):
                    ws.cell(row=i, column=col).fill = fill

        ws.sheet_view.showGridLines = False
        ws.print_area = f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}"

    # ✅ MIMO smyčku!
    file_path = "smenar_2026.xlsx"
    wb.save(file_path)

    return send_file(file_path, as_attachment=True)


@app.route("/dashboard")
def dashboard():
    selected_date_param = request.args.get("date")

    if selected_date_param:
        selected_date_obj = datetime.strptime(selected_date_param, "%Y-%m-%d")
    else:
        selected_date_obj = datetime(2026, 1, 1)

    selected_date = selected_date_obj.strftime("%d.%m.%Y")
    selected_date_iso = selected_date_obj.strftime("%Y-%m-%d")

    data = []

    try:
        with open("smenar_actual.csv", encoding="utf-8") as f:
            reader = csv.DictReader(f, delimiter=";")
            all_data = list(reader)
    except FileNotFoundError:
        all_data = []

    for row in all_data:
        try:
            row_raw = row["Datum"].strip()
            parsed = None
            for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
                try:
                    parsed = datetime.strptime(row_raw, fmt)
                    break
                except:
                    continue
            if not parsed:
                continue
            if parsed.date() == selected_date_obj.date():
                row["date_obj"] = parsed
                row["has_manual_minutes"] = row.get("ManualMinutes", "0") == "1"
                data.append(row)
        except:
            continue

    return render_template(
        "dashboard.html",
        data=data,
        selected_date=selected_date,
        selected_date_iso=selected_date_iso,
        get_shift_time=get_shift_time
    )


if __name__ == "__main__":
    app.run(debug=True)